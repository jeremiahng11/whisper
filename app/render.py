"""Output formats: Markdown minutes, plain text, SRT subtitles, JSON."""
import re
from datetime import datetime
from typing import Optional


def ts(sec: float) -> str:
    sec = int(sec)
    h, m, s = sec // 3600, sec % 3600 // 60, sec % 60
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def srt_ts(sec: float) -> str:
    ms = int(round(sec * 1000))
    h, ms = divmod(ms, 3600000)
    m, ms = divmod(ms, 60000)
    s, ms = divmod(ms, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def recorded_at(filename: str, fallback: float) -> datetime:
    """Idea Saver names files like 2026-10-08_1535.wav - use that as the recording time."""
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})[_ T-]?(\d{2})[:\-]?(\d{2})", filename or "")
    if m:
        try:
            return datetime(*map(int, m.groups()))
        except ValueError:
            pass
    return datetime.fromtimestamp(fallback)


def paragraphs(segments: list[dict], max_len: int = 600, gap: float = 2.5) -> list[dict]:
    """Merge segments into readable paragraphs: new paragraph on speaker change, long pause or length."""
    out: list[dict] = []
    for s in segments:
        p = out[-1] if out else None
        if (p and p["speaker"] == s.get("speaker", "") and s["start"] - p["end"] < gap
                and len(p["text"]) + len(s["text"]) < max_len):
            p["text"] += " " + s["text"].strip()
            p["end"] = s["end"]
        else:
            out.append({"start": s["start"], "end": s["end"], "speaker": s.get("speaker", ""), "text": s["text"].strip()})
    return out


# ---------------------------------------------------------------- speaker names
def names_of(job: dict) -> dict:
    return {k: v for k, v in ((job.get("options") or {}).get("speaker_names") or {}).items() if v}


def apply_names(tr: dict, names: dict) -> dict:
    """Labels -> names; a line with two people talking at once reads "Alan + Wei Ling"."""
    if not names and not any(s.get("also") for s in tr.get("segments", [])):
        return tr
    out = dict(tr)

    def who(s):
        main = names.get(s.get("speaker", ""), s.get("speaker", ""))
        also = [names.get(a, a) for a in (s.get("also") or []) if a and a != s.get("speaker")]
        return " + ".join([main] + also) if main and also else main

    out["segments"] = [{**s, "speaker": who(s)} for s in tr.get("segments", [])]
    return out


def rename_text(text: str, names: dict, at_summary: Optional[dict] = None) -> str:
    """Summary was written with the names known then (at_summary); bring it up to date with `names`."""
    if not text:
        return text
    at_summary = at_summary or {}
    for label in sorted(set(names) | set(at_summary), key=len, reverse=True):
        new = names.get(label) or label
        old = at_summary.get(label) or label
        if old != new:
            text = re.sub(rf"(?<!\w){re.escape(old)}(?!\w)", new, text)
    return text


# ---------------------------------------------------------------- flagged moments (mark key on the device)
_MARK = re.compile(r"^\s*\[(?:(\d+):)?(\d{1,3}):(\d{2})\]\s*(?:!!|\*\*?MARK\*?\*?|\u2605)\s*(.*)$", re.I)


def parse_marks(notes: str) -> list[tuple[float, str]]:
    out = []
    for line in (notes or "").splitlines():
        m = _MARK.match(line)
        if m:
            h, mi, se, txt = m.groups()
            out.append((int(h or 0) * 3600 + int(mi) * 60 + int(se), txt.strip()))
    return out


def around(tr: dict, t: float, before: float = 40, after: float = 20) -> str:
    return " ".join(
        (f"{s['speaker']}: " if s.get("speaker") else "") + s["text"].strip()
        for s in tr.get("segments", []) if s["end"] >= t - before and s["start"] <= t + after)


def flagged_text(tr: dict, marks: list[tuple[float, str]]) -> str:
    return "\n\n".join(f"[{ts(t)}] {('(' + note + ') ') if note else ''}{around(tr, t)[:1500]}" for t, note in marks)


def markdown(job: dict, tr: dict, summary_md: Optional[str], title: str = "") -> str:
    opts = job.get("options") or {}
    names = names_of(job)
    tr = apply_names(tr, names)
    summary_md = rename_text(summary_md or "", names, opts.get("names_at_summary")) or None
    when = recorded_at(job["filename"], job["created_at"])
    mins = tr.get("duration", 0) / 60
    lines = [f"# {title or job.get('title') or job['filename']}", ""]
    meta = [when.strftime("%Y-%m-%d %H:%M"), f"{mins:.0f} min" if mins >= 1 else f"{tr.get('duration', 0):.0f} s"]
    if tr.get("language"):
        meta.append(tr["language"])
    meta.append(f"whisper {tr.get('model', '')}")
    lines += ["*" + " · ".join(meta) + "*", ""]
    if summary_md:
        lines += [summary_md.strip(), ""]
    elif job.get("summary_error"):
        lines += [f"> Summary not available: {job['summary_error']}", ""]
    marks = parse_marks(opts.get("notes", ""))
    if marks:
        lines += ["## Flagged moments", ""]
        for t, note in marks:
            ctx = around(tr, t, 25, 15)
            ctx = (ctx[:280] + "...") if len(ctx) > 280 else ctx
            lines += [f"- `[{ts(t)}]` " + (f"**{note}** - " if note else "") + (f"*{ctx}*" if ctx else ""), ""]
    notes = opts.get("notes", "").strip()
    if notes:
        lines += ["## My notes", "", notes, ""]
    lines += ["## Transcript", ""]
    for p in paragraphs(tr.get("segments", [])):
        who = f"**{p['speaker']}:** " if p["speaker"] else ""
        lines += [f"`[{ts(p['start'])}]` {who}{p['text']}", ""]
    if not tr.get("segments"):
        lines += ["*(no speech found)*", ""]
    return "\n".join(lines).rstrip() + "\n"


def text(tr: dict) -> str:
    out = []
    for p in paragraphs(tr.get("segments", [])):
        who = f"{p['speaker']}: " if p["speaker"] else ""
        out.append(f"[{ts(p['start'])}] {who}{p['text']}")
    return "\n\n".join(out) + "\n"


def srt(tr: dict) -> str:
    out = []
    for i, s in enumerate(tr.get("segments", []), 1):
        who = f"[{s['speaker']}] " if s.get("speaker") else ""
        out.append(f"{i}\n{srt_ts(s['start'])} --> {srt_ts(s['end'])}\n{who}{s['text'].strip()}\n")
    return "\n".join(out)


def transcript_for_llm(tr: dict) -> str:
    """Compact transcript fed to the summariser (one line per paragraph, with time and speaker)."""
    return "\n".join(
        f"[{ts(p['start'])}] {p['speaker'] + ': ' if p['speaker'] else ''}{p['text']}"
        for p in paragraphs(tr.get("segments", []))
    )


# ---------------------------------------------------------------- Word export
def _runs(par, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            from docx.shared import Pt, RGBColor
            r.font.size = Pt(9)                            # timestamps: small and grey, like the web page
            r.font.color.rgb = RGBColor.from_string("5D6B7A")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


# the web page's speaker colours (light theme), in speaker-number order
SPEAKER_COLOURS = ["2F6FD6", "E07A2E", "23935F", "B23F78", "7354CC", "B8901A", "1F9DB3", "8A6545"]


def _colour_names(par, colours: dict) -> None:
    """Bold "Name:" (or "Name + Other:") runs get that speaker's colour."""
    from docx.shared import RGBColor
    for r in par.runs:
        if r.bold:
            who = r.text.rstrip(": ").split(" + ")[0].strip().lower()
            if who in colours:
                r.font.color.rgb = RGBColor.from_string(colours[who])


def docx(md: str, speakers: Optional[list] = None) -> bytes:
    """speakers = [(name, colour hex, share 0..1)]: names in the transcript get their colour, and a line under the
    title shows who spoke how much (like the web page's voice ribbon)."""
    import io
    from docx import Document
    from docx.shared import Pt, RGBColor
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    colours = {n.lower(): c for n, c, _ in (speakers or [])}
    legend_done = not speakers
    for line in md.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            d.add_heading(m.group(2).strip(), level=len(m.group(1)) - 1)
            if not legend_done and len(m.group(1)) == 1:      # under the title: who spoke how much
                legend_done = True
                p = d.add_paragraph()
                for i, (n, c, share) in enumerate(speakers):
                    if i:
                        p.add_run("    ")
                    sq = p.add_run("\u25a0 ")
                    sq.font.color.rgb = RGBColor.from_string(c)
                    nm = p.add_run(n)
                    nm.bold = True
                    nm.font.color.rgb = RGBColor.from_string(c)
                    pc = p.add_run(f" {round(share * 100)}%")
                    pc.font.size = Pt(9)
                    pc.font.color.rgb = RGBColor.from_string("5D6B7A")
            continue
        m = re.match(r"^\s*[-*]\s+\[( |x|X)\]\s+(.*)$", s)
        if m:
            p = d.add_paragraph(style="List Bullet")
            p.add_run("\u2611 " if m.group(1).strip() else "\u2610 ")
            _runs(p, m.group(2))
            continue
        m = re.match(r"^\s*[-*]\s+(.*)$", s)
        if m:
            _runs(d.add_paragraph(style="List Bullet"), m.group(1))
            continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)$", s)
        if m:
            _runs(d.add_paragraph(style="List Number"), m.group(1))
            continue
        if s.startswith(">"):
            p = d.add_paragraph()
            _runs(p, s.lstrip("> "))
            for r in p.runs:
                r.italic = True
            continue
        p = d.add_paragraph()
        _runs(p, s)
        if colours:
            _colour_names(p, colours)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
