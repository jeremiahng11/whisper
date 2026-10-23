"""Your own sources for the knowledge base: web pages, product sheets, price lists, FAQs.

Add one on Account > Watch AI:
  * a web page link - optionally following links on the same site (`crawl`: up to that many more pages under the
    same address, e.g. https://example.com/products/ and every product page below it);
  * a file: PDF, Word (.docx), text, Markdown or HTML;
  * pasted text.
The server reads it (a web page: fetched again once a day; only re-read by the AI when its text changed) and the AI
turns it into facts like the meetings' ("How much is the Round watch?" -> "SGD 299, with a one-year warranty.").

Kept in DATA_DIR/knowledge/sources/<id>.json (+ <id>.txt, the text it read).
"""
import hashlib
import html
import io
import json
import os
import logging
import re
import threading
import time
import uuid
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urldefrag, urljoin, urlparse

import httpx

from . import knowledge, summarize
from .config import settings

log = logging.getLogger("sources")
MAX_TEXT = 400_000              # per source (about 100 pages of text)
CHUNK = 9000                    # characters the AI reads at a time
PER_CHUNK = 20                  # facts per chunk
MAX_CRAWL = 60
REFETCH_S = 86400


def sdir() -> Path:
    d = knowledge.kdir() / "sources"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _meta_path(sid: str) -> Path:
    return sdir() / f"{sid}.json"


def _ok_id(sid: str) -> bool:
    return bool(re.fullmatch(r"[0-9a-f]{12}", sid or ""))


def get(sid: str) -> dict | None:
    if not _ok_id(sid):
        return None
    try:
        return json.loads(_meta_path(sid).read_text())
    except (OSError, ValueError):
        return None


_locks: dict[str, threading.Lock] = {}
_locks_guard = threading.Lock()


def _lock(sid: str) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(sid, threading.Lock())


def _save(m: dict) -> None:
    tmp = _meta_path(m["id"]).with_suffix(f".{threading.get_ident()}.tmp")
    tmp.write_text(json.dumps(m, ensure_ascii=False))
    tmp.replace(_meta_path(m["id"]))


def listing() -> list[dict]:
    out = []
    for p in sorted(sdir().glob("*.json")):
        if p.name.endswith(".news.json"):           # a news/research source's article store, not a source
            continue
        try:
            m = json.loads(p.read_text())
        except ValueError:
            continue
        out.append({k: v for k, v in m.items() if k != "facts"} | {"facts": len(m.get("facts") or [])})
    out.sort(key=lambda m: m.get("added", 0), reverse=True)
    n = 0                                           # the order the background reader takes them in (one at a time)
    for m in out:
        if m.get("state") in ("new", "changed"):
            n += 1
            m["queue"] = n
    return out


def all_facts() -> list[dict]:
    fs = []
    for p in sdir().glob("*.json"):
        if p.name.endswith(".news.json"):
            continue
        try:
            fs += json.loads(p.read_text()).get("facts") or []
        except ValueError:
            pass
    return fs


# ---------------------------------------------------------------- text from HTML, PDF, Word
class _Text(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form", "iframe", "template"}
    BLOCK = {"p", "div", "li", "tr", "br", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "td", "th", "dd", "dt"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.skip, self.title, self.in_title, self.links = [], 0, "", False, []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self.in_title = True
        if tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self.in_title = False
        if tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip:
            self.out.append(data)


def html_text(raw: str) -> tuple[str, str, list[str]]:
    p = _Text()
    try:
        p.feed(raw)
    except Exception:                               # noqa: BLE001 - broken HTML: keep what was read
        pass
    text = re.sub(r"[ \t\r\f\v]+", " ", html.unescape("".join(p.out)))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return p.title.strip(), text, p.links


def file_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ValueError("PDF needs the pypdf package on the server (it's in requirements.txt from 1.12.1)")
        r = PdfReader(io.BytesIO(data))
        return "\n\n".join((pg.extract_text() or "") for pg in r.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    raw = data.decode("utf-8", "replace")
    if ext in (".html", ".htm"):
        return html_text(raw)[1]
    if ext in (".txt", ".md", ".csv", ".json", ""):
        return raw
    raise ValueError(f"can't read {ext} files - use PDF, Word (.docx), text, Markdown, CSV or HTML")


# ---------------------------------------------------------------- web pages
# Many sites turn away clients that don't look like a browser (HTTP 403 from Zendesk, Cloudflare ...): we ask like one.
BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/128.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,application/json;q=0.8,*/*;q=0.7",
    "Accept-Language": "en-US,en;q=0.9",
}
MIN_PAGE_TEXT = 200                                 # less than this from plain HTML: the page is built by JavaScript


def _get(url: str, timeout: float = 25) -> httpx.Response:
    return httpx.get(url, timeout=timeout, follow_redirects=True, headers=BROWSER_HEADERS)


def browser_available() -> bool:
    try:
        import playwright.sync_api  # noqa: F401
        return True
    except ImportError:
        return False


def render(url: str) -> tuple[str, str, list[str]]:
    """Open the page in a headless Chromium (sites that build their pages with JavaScript, or that only answer
    browsers) and read what's on screen: (title, text, links)."""
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch(args=["--no-sandbox", "--disable-dev-shm-usage"])
        try:
            page = b.new_page(user_agent=BROWSER_HEADERS["User-Agent"], locale="en-US")
            page.goto(url, wait_until="networkidle", timeout=45000)
            page.wait_for_timeout(1500)               # late content (animations, lazy lists)
            html_ = page.content()
            links = page.eval_on_selector_all("a[href]", "els => els.map(e => e.href)")
        finally:
            b.close()
    title, text, more = html_text(html_)
    return title, text, list(dict.fromkeys(links + more))


def zendesk_articles(url: str, limit: int = 2000) -> list[dict]:
    """A Zendesk Help Center (/hc/<locale>): every article through its public API, which is how such sites hand out
    their text (the pages themselves often refuse anything but a browser). [] if it isn't one."""
    u = urlparse(url)
    m = re.match(r"/hc/([A-Za-z-]+)", u.path or "")
    if not m:
        return []
    api = f"{u.scheme}://{u.netloc}/api/v2/help_center/{m.group(1)}/articles.json?per_page=100&page=1"
    out = []
    while api and len(out) < limit:
        try:
            r = _get(api)
        except httpx.HTTPError:
            break
        if r.status_code != 200 or "json" not in r.headers.get("content-type", ""):
            break
        d = r.json()
        for a in d.get("articles") or []:
            if a.get("draft"):
                continue
            body = html_text(a.get("body") or "")[1]
            if body:
                out.append({"url": a.get("html_url") or url, "title": a.get("title") or "", "text": body})
        api = d.get("next_page")
        time.sleep(0.3)
    return out


def _page(u: str, use_browser: bool) -> tuple[str, str, list[str], str]:
    """One page: (title, text, links, how) - plain HTTP first, the headless browser when that gets nothing useful."""
    err = ""
    try:
        r = _get(u)
        ctype = r.headers.get("content-type", "")
        if r.status_code == 200 and "pdf" in ctype:
            return Path(urlparse(u).path).name, file_text("x.pdf", r.content), [], "pdf"
        if r.status_code == 200 and ("html" in ctype or "text" in ctype):
            title, text, links = html_text(r.text)
            if len(text) >= MIN_PAGE_TEXT or not use_browser:
                return title, text, links, "http"
        else:
            err = f"answered HTTP {r.status_code}"
    except httpx.HTTPError as e:
        err = str(e)
    if use_browser:
        try:
            title, text, links = render(u)
            return title, text, links, "browser"
        except Exception as e:                       # noqa: BLE001
            err = (err + "; " if err else "") + f"browser: {str(e).splitlines()[0][:160]}"
    raise ValueError(err or "no text")


def fetch_site(url: str, crawl: int = 0) -> list[dict]:
    """The page, and up to `crawl` more pages under the same address: [{"url", "title", "text"}]."""
    start = urldefrag(url)[0]
    help_center = zendesk_articles(start, 2000 if crawl else 200)
    if help_center:
        return help_center
    base = urlparse(start)
    prefix = start.rsplit("/", 1)[0] + "/" if not start.endswith("/") else start
    todo, seen, pages = [start], {start}, []
    use_browser = browser_available()
    while todo and len(pages) <= min(crawl, MAX_CRAWL):
        u = todo.pop(0)
        try:
            title, text, links, how = _page(u, use_browser)
        except ValueError as e:
            if not pages:
                hint = "" if use_browser else " (this server has no headless browser: rebuild with INSTALL_BROWSER=1)"
                raise ValueError(f"couldn't read {u}: {e}{hint}")
            continue
        pages.append({"url": u, "title": title or u, "text": text, "how": how})
        for h in links:
            v = urldefrag(urljoin(u, h))[0]
            pv = urlparse(v)
            if pv.scheme in ("http", "https") and pv.netloc == base.netloc and v.startswith(prefix) and v not in seen \
                    and not re.search(r"\.(jpg|jpeg|png|gif|svg|webp|zip|mp4|mp3|css|js)(\?|$)", v, re.I):
                seen.add(v)
                todo.append(v)
        if crawl:
            time.sleep(0.5)                         # be gentle with the site
    return pages


# ---------------------------------------------------------------- news about a topic (research in the background)
# A "news" source follows a topic ("Aleta Planet", a competitor, your industry): once a day the server searches the
# news (Bing News, else Google News - both free, no key), reads the new articles (the headless browser when a site
# needs it) and keeps the last 60 days. The AI turns them into dated facts like any other source.
NEWS_KEEP_DAYS = 60
NEWS_MAX = 60
NEWS_NEW_PER_RUN = 12

NEWS_PROMPT = """From these news articles about "{topic}", write the facts someone in the company may later ask about:
what happened, who, when, numbers, partnerships, products, regulation. Return ONLY a JSON array (at most {n} items) of
objects with keys "q" (a short spoken question, 3-8 words), "a" (one short sentence, at most 30 words, that says
WHEN - the article date - and the source, e.g. "On 12 October 2026 (The Business Times) ..."), "kind" (use "news"),
"topic" (1-4 words), "alts" (2-3 other wordings). Only what the articles clearly say.

Articles:
\"\"\"
{text}
\"\"\""""


RESEARCH_PROMPT = """These are web pages and news found by searching for "{topic}". {asked}Write the facts someone in the
company may later ask about {topic}: what it is, products and services, prices, places, dates, people and their roles,
partnerships (e.g. licensed characters), latest launches and news. Ignore pages that are about something else with a
similar name. Return ONLY a JSON array (at most {n} items) of objects with keys "q" (a short spoken question, 3-8
words), "a" (one short sentence, at most 30 words; for news say WHEN and the source), "kind" (one of product, price,
policy, person, date, fact, news), "topic" (1-4 words), "alts" (2-3 other wordings). Only what the pages clearly say.

Pages:
\"\"\"
{text}
\"\"\""""


def _rss_items(xml_text: str) -> list[dict]:
    import xml.etree.ElementTree as ET
    from email.utils import parsedate_to_datetime
    from urllib.parse import parse_qs
    out = []
    try:
        root = ET.fromstring(xml_text)
    except ET.ParseError:
        return out
    for it in root.iter("item"):
        g = lambda tag: (it.findtext(tag) or "").strip()   # noqa: E731
        link = g("link")
        q = parse_qs(urlparse(link).query)
        if "bing.com" in link and q.get("url"):          # Bing wraps the article's address
            link = q["url"][0]
        try:
            when = parsedate_to_datetime(g("pubDate")).strftime("%Y-%m-%d")
        except (TypeError, ValueError):
            when = time.strftime("%Y-%m-%d")
        src = it.find("source")
        out.append({"url": link, "title": html.unescape(g("title")), "date": when,
                    "source": (src.text or "").strip() if src is not None else urlparse(link).netloc,
                    "snippet": html_text(g("description"))[1][:600]})
    return out


def news_search(topic: str, timeout: float = 25) -> list[dict]:
    from urllib.parse import quote_plus
    for feed in (f"https://www.bing.com/news/search?q={quote_plus(topic)}&format=rss&setlang=en",
                 f"https://news.google.com/rss/search?q={quote_plus(topic)}&hl=en-SG&gl=SG&ceid=SG:en"):
        try:
            r = _get(feed, timeout)
        except httpx.HTTPError:
            continue
        if r.status_code == 200:
            items = _rss_items(r.text)
            if items:
                return items
    return []


search_status: dict = {"provider": "", "error": "", "at": 0.0}


def _tavily(query: str, timeout: float) -> list[dict]:
    r = httpx.post("https://api.tavily.com/search", timeout=timeout,
                   headers={"Authorization": f"Bearer {settings.tavily_api_key}"},
                   json={"query": query, "max_results": 6, "search_depth": "basic", "include_answer": False})
    r.raise_for_status()
    return [{"url": x.get("url", ""), "title": x.get("title", ""), "snippet": (x.get("content") or "")[:500],
             "date": time.strftime("%Y-%m-%d"), "source": urlparse(x.get("url", "")).netloc}
            for x in (r.json().get("results") or []) if x.get("url")]


def _brave(query: str, timeout: float) -> list[dict]:
    from urllib.parse import quote_plus
    r = httpx.get(f"https://api.search.brave.com/res/v1/web/search?count=6&q={quote_plus(query)}", timeout=timeout,
                  headers={"X-Subscription-Token": settings.brave_api_key, "Accept": "application/json"})
    r.raise_for_status()
    return [{"url": x.get("url", ""), "title": html_text(x.get("title", ""))[1],
             "snippet": html_text(x.get("description", ""))[1][:500], "date": time.strftime("%Y-%m-%d"),
             "source": urlparse(x.get("url", "")).netloc}
            for x in ((r.json().get("web") or {}).get("results") or []) if x.get("url")]


def web_search(query: str, timeout: float = 25) -> list[dict]:
    """Web results for a query: Tavily / Brave when their key is set (reliable), else Bing's RSS results, else
    DuckDuckGo's HTML page (free, no key): [{"url", "title", "snippet", "date", "source"}]."""
    # SEARCH_ORDER (default "brave,tavily"): the first with a key; it fails -> the next, then the free ones
    known = {"brave": ("Brave", settings.brave_api_key, _brave), "tavily": ("Tavily", settings.tavily_api_key, _tavily)}
    order = [known[x.strip().lower()] for x in (os.getenv("SEARCH_ORDER") or "brave,tavily").split(",") if x.strip().lower() in known]
    failed = []
    for name, key, fn in order:
        if not key:
            continue
        try:
            out = fn(query, timeout)
            search_status.update(provider=name, error="", fell_back="; ".join(failed), at=time.time())
            return out
        except Exception as e:                       # noqa: BLE001 - the next one, then the free ones
            failed.append(f"{name}: {str(e)[:120]}")
            search_status.update(provider=name, error=failed[-1], fell_back="", at=time.time())
            log.warning("web search %s failed: %s", name, e)
    out = _free_search(query, timeout)
    if not (settings.tavily_api_key or settings.brave_api_key):
        search_status.update(provider="Bing/DuckDuckGo", error="" if out else "no results (often blocked from servers)",
                             at=time.time())
    return out


def _free_search(query: str, timeout: float) -> list[dict]:
    from urllib.parse import parse_qs, quote_plus, unquote
    try:
        r = _get(f"https://www.bing.com/search?format=rss&setlang=en&q={quote_plus(query)}", timeout)
        if r.status_code == 200:
            items = _rss_items(r.text)
            if items:
                return items
    except httpx.HTTPError:
        pass
    try:
        r = _get(f"https://html.duckduckgo.com/html/?q={quote_plus(query)}", timeout)
    except httpx.HTTPError:
        return []
    if r.status_code != 200:
        return []
    out = []
    for m in re.finditer(r'<a[^>]+class="result__a"[^>]+href="([^"]+)"[^>]*>(.*?)</a>(.*?)(?=<a[^>]+class="result__a"|$)', r.text, re.S):
        href = html.unescape(m.group(1))
        q = parse_qs(urlparse(href).query)
        url = unquote(q["uddg"][0]) if q.get("uddg") else href
        if not url.startswith("http"):
            continue
        snip = re.search(r'class="result__snippet"[^>]*>(.*?)</a>', m.group(3), re.S)
        out.append({"url": url, "title": html_text(m.group(2))[1], "snippet": html_text(snip.group(1))[1] if snip else "",
                    "date": time.strftime("%Y-%m-%d"), "source": urlparse(url).netloc})
        if len(out) >= 15:
            break
    return out


def _news_path(sid: str) -> Path:
    return sdir() / f"{sid}.news.json"


def refresh_news(m: dict) -> str:
    """New articles about the topic into the source's store; returns the text the AI reads (newest first)."""
    try:
        store = json.loads(_news_path(m["id"]).read_text())
    except (OSError, ValueError):
        store = {}
    topic = m["query"]
    research = m["kind"] == "research"
    words = sorted(knowledge._words(m.get("must") or topic)) if research else \
        [w for w in re.findall(r"\w+", topic.lower()) if len(w) > 2]
    use_browser = browser_available()
    new = 0
    search = m.get("search") or topic               # (a question is searched with the company's name added)
    found = news_search(search)
    if research:                                    # the web as well as the news; the web first (more lasting)
        seen_u = set()
        found = [x for x in web_search(search) + found if not (x["url"] in seen_u or seen_u.add(x["url"]))]
    m["searched"] = len(found)                       # (the page says "the search found nothing" vs "nothing about it")
    for it in found:
        if it["url"] in store or new >= NEWS_NEW_PER_RUN:
            continue
        hay = (it["title"] + " " + it["snippet"] + " " + it["url"]).lower()
        # a research topic: all its words (one may be missing from a long one) - "Sanrio Aleta Adventure" isn't any
        # page that mentions Sanrio
        need = m.get("need") or ((len(words) if len(words) <= 3 else len(words) - 1) if research else 1)
        if words and sum(w in hay for w in words) < need:   # search engines pad with loosely related results
            continue
        text = it["snippet"]
        try:
            _, body, _, _ = _page(it["url"], use_browser)
            if len(body) > len(text):
                text = body[:6000]
        except ValueError:
            pass
        store[it["url"]] = {**{k: it[k] for k in ("title", "date", "source")}, "text": text}
        new += 1
        time.sleep(0.5)
    days = 365 if research else NEWS_KEEP_DAYS      # web pages (products, profiles) stay useful longer than news
    cutoff = time.strftime("%Y-%m-%d", time.localtime(time.time() - days * 86400))
    keep = sorted(((u, a) for u, a in store.items() if a["date"] >= cutoff), key=lambda x: x[1]["date"], reverse=True)[:NEWS_MAX]
    store = dict(keep)
    _news_path(m["id"]).write_text(json.dumps(store, ensure_ascii=False))
    m["articles"] = len(store)
    m["new_articles"] = new
    return "\n\n".join(f"## {a['title']} ({a['source']}, {a['date']})\n{u}\n\n{a['text']}" for u, a in keep)


# ---------------------------------------------------------------- add / learn / remove
def add(kind: str, title: str = "", url: str = "", text: str = "", filename: str = "", data: bytes = b"",
        crawl: int = 0) -> dict:
    sid = uuid.uuid4().hex[:12]
    m = {"id": sid, "kind": kind, "title": title, "added": time.time(), "fetched": 0, "sig": "", "facts": [], "error": "",
         "state": "new"}
    if kind == "url":
        u = urlparse(url.strip())
        if u.scheme not in ("http", "https") or not u.netloc:
            raise ValueError("give a full web address (https://...)")
        m.update(url=url.strip(), crawl=max(0, min(int(crawl or 0), MAX_CRAWL)), title=title or u.netloc + u.path)
    elif kind in ("news", "research"):
        topic = re.sub(r"\s+", " ", (url or title or text or "").strip())[:120]
        if len(topic) < 3:
            raise ValueError("say what to follow (a company, a person, a topic)")
        m.update(query=topic, title=title or (f"News: {topic}" if kind == "news" else f"Research: {topic}"))
    elif kind == "file":
        body = file_text(filename, data)
        if len(body.strip()) < 20:
            raise ValueError("no text found in that file (a scanned PDF has none)")
        (sdir() / f"{sid}.txt").write_text(body[:MAX_TEXT])
        m.update(title=title or filename, filename=filename)
    elif kind == "text":
        if len(text.strip()) < 20:
            raise ValueError("paste a bit more text")
        (sdir() / f"{sid}.txt").write_text(text[:MAX_TEXT])
        m.update(title=title or text.strip().split("\n")[0][:60])
    else:
        raise ValueError("kind must be url, file or text")
    _save(m)
    return m


def reset_stale() -> None:
    """At start-up: a source left "reading" by a restart is read again."""
    for p in sdir().glob("*.json"):
        if p.name.endswith(".news.json"):
            continue
        try:
            m = json.loads(p.read_text())
            if m.get("state") == "reading":
                m.update(state="changed", progress="")
                _save(m)
        except (OSError, ValueError):
            pass


def remove(sid: str) -> bool:
    if not get(sid):
        return False
    for ext in (".json", ".txt", ".news.json"):
        try:
            (sdir() / f"{sid}{ext}").unlink()
        except OSError:
            pass
    knowledge.rebuild()
    return True


def due(m: dict) -> bool:
    if m.get("state") in ("new", "changed"):
        return True
    if m.get("state") == "reading" and time.time() - (m.get("since") or 0) > 3600 and not _lock(m["id"]).locked():
        return True                                 # the server restarted while reading it: start again
    if m.get("once") and m.get("fetched"):
        return False                                # a one-off research (a question you asked): done
    every = 7 * REFETCH_S if m.get("kind") == "research" else REFETCH_S   # web research weekly, news daily
    return m.get("kind") in ("url", "news", "research") and time.time() - (m.get("fetched") or 0) > every


def _chunks(text: str) -> list[str]:
    parts, cur = [], ""
    for para in re.split(r"\n\s*\n", text):
        if len(cur) + len(para) > CHUNK and cur:
            parts.append(cur)
            cur = ""
        cur += para[: CHUNK * 2] + "\n\n"
    if cur.strip():
        parts.append(cur)
    return parts


def learn(sid: str, force: bool = False) -> int:
    """(Re)read a source; the AI only reads it again if its text changed. Returns how many facts it has."""
    with _lock(sid):                                # one reader per source (the button and the sweep may meet)
        return _learn(sid, force)


def _learn(sid: str, force: bool) -> int:
    m = get(sid)
    if not m:
        return 0
    m.update(state="reading", since=time.time(), progress="fetching" if m["kind"] == "url" else "")
    _save(m)
    try:
        if m["kind"] == "url":
            pages = fetch_site(m["url"], m.get("crawl", 0))
            text = "\n\n".join(f"## {p['title']}\n{p['url']}\n\n{p['text']}" for p in pages)[:MAX_TEXT]
            if sum(len(p["text"]) for p in pages) < 150:
                raise ValueError("that page has almost no text for the server to read even in a browser (a login page?) "
                                 "- add a PDF or paste the text" if browser_available() else
                                 "that page has almost no text for the server to read (the site builds it in the "
                                 "browser) - rebuild Oracle Minutes with INSTALL_BROWSER=1 (headless Chromium), or add "
                                 "a PDF / paste the text")
            (sdir() / f"{sid}.txt").write_text(text)
            m["pages"] = len(pages)
            if pages and (not m.get("title") or m["title"].startswith(urlparse(m["url"]).netloc)):
                m["title"] = pages[0]["title"][:80] or m["title"]
        elif m["kind"] in ("news", "research"):
            m["progress"] = "searching the news" if m["kind"] == "news" else "searching the web and the news"
            _save(m)
            text = refresh_news(m)[:MAX_TEXT]
            (sdir() / f"{sid}.txt").write_text(text)
            if not text:
                m.update(fetched=time.time(), state="ok", error="", progress="", facts=[])
                _save(m)
                knowledge.rebuild()
                return 0
        else:
            text = (sdir() / f"{sid}.txt").read_text()
        m["fetched"] = time.time()
        m["chars"] = len(text)
        sig = hashlib.sha1(text.encode()).hexdigest()[:12]
        if sig == m.get("sig") and not force:
            m.update(state="ok", error="")
            _save(m)
            return len(m.get("facts") or [])
        date = time.strftime("%Y-%m-%d")
        facts, seen = [], set()
        parts = _chunks(text)
        for i, part in enumerate(parts):
            m["progress"] = f"part {i + 1} of {len(parts)}"   # (the AI may be busy with your meetings first)
            _save(m)
            if m["kind"] == "news":
                prompt = NEWS_PROMPT.format(topic=m["query"], n=PER_CHUNK, text=part)
            elif m["kind"] == "research":
                prompt = RESEARCH_PROMPT.format(topic=m["query"], n=PER_CHUNK, text=part,
                                                asked=(f'Someone asked: "{m["question"]}" - if these pages answer it, '
                                                       f'make that the first fact, worded as they asked.\n') if m.get("question") else "")
            else:
                prompt = knowledge.DOC_PROMPT.format(n=PER_CHUNK, title=m["title"], where=m.get("url") or m.get("filename") or "pasted",
                                                     text=part)
            for f in knowledge.parse_facts(summarize.chat([{"role": "system", "content": "You extract facts from documents as JSON."},
                                                           {"role": "user", "content": prompt}])):
                k = knowledge._norm(f["q"])
                if k in seen:
                    continue
                seen.add(k)
                if m["kind"] == "news":
                    f["kind"] = "news"
                f.update(job="", source=sid, date=date, meeting=m["title"])
                facts.append(f)
        m.update(facts=facts, sig=sig, state="ok", error="", progress="")
    except Exception as e:                          # noqa: BLE001
        m.update(state="error", error=summarize.explain_llm_error(e) if not isinstance(e, ValueError) else str(e))
        _save(m)
        raise
    _save(m)
    knowledge.rebuild()
    return len(facts)
